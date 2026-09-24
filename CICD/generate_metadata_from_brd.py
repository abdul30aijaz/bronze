#!/usr/bin/env python3
"""
Generate MDIF metadata/configs/<job_name>/<table_name>.json file(s) from one or more
BRD workbooks (TECH DOC + STTM sheet pair), following the same layout as every
BRD produced by the team (see AMS_BRD.xlsx as the reference example).

Usage:
    python CICD/generate_metadata_from_brd.py <brd1.xlsx> [<brd2.xlsx> ...] \
        --repo-root . \
        --notify-emails a@company.com,b@company.com \
        --file-format excel \
        --summary-out /tmp/brd_summary.json

Writes:
    <repo-root>/metadata/configs/<job_name>/<table_name>.json   (one per BRD)

Also writes a JSON summary (list of per-file results, including any assumptions /
warnings a human reviewer should double check) to --summary-out, which the CI
workflow turns into the pull request description.

This script only depends on openpyxl (already used elsewhere in this repo).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Conventions pulled from the existing metadata/configs/**/*.json files.
# Keep these in sync with utils/schema_config.py's supported type keys.
# --------------------------------------------------------------------------- #

_STRING_TYPES = {"string", "varchar", "char", "text"}
_DOUBLE_TYPES = {
    "int", "integer", "smallint", "tinyint", "bigint",
    "float", "double", "decimal", "numeric", "number",
}
_DATE_TYPES = {"date"}
_TIMESTAMP_TYPES = {"timestamp", "datetime"}

# Column-name substrings that are flagged as PII candidates for reviewer sign-off.
# Populated into `pii_columns` (masking itself stays disabled per masking_utils.py
# until a secret store is wired up) but MUST be eyeballed in the PR review.
_PII_HINTS = (
    "email", "phone", "ssn", "social_security", "address", "dob",
    "date_of_birth", "birth_date", "passport", "credit_card", "card_number",
    "ip_address", "first_name", "last_name", "full_name", "customer_name",
    "account_number", "national_id",
)

# REGION -> short code used in job-folder names (na, eu, ...). Extend as new
# regions show up in BRDs; unknown regions fall back to an acronym of the words.
_REGION_CODES = {
    "NORTH AMERICA": "na",
    "SOUTH AMERICA": "sa",
    "LATIN AMERICA": "latam",
    "EUROPE": "eu",
    "ASIA PACIFIC": "apac",
    "APAC": "apac",
    "MIDDLE EAST": "me",
    "AFRICA": "af",
    "GLOBAL": "global",
}

_TECH_DOC_SHEET_HINTS = ("TECH DOC", "TECHNICAL")
_STTM_SHEET_HINTS = ("STTM", "SOURCE TO TARGET", "MAPPING")


class BRDFormatError(ValueError):
    """Raised when a workbook doesn't match the expected BRD layout."""


@dataclass
class BRDResult:
    brd_file: str
    job_name: str
    table_name: str
    output_path: str
    region: str
    market: str
    domain: str
    source_name: str
    cadence: str
    load_type: str
    primary_key: list
    column_count: int
    file_format: str
    pii_candidates: list
    warnings: list = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Sheet lookup helpers
# --------------------------------------------------------------------------- #

def _find_sheet(wb, hints) -> "openpyxl.worksheet.worksheet.Worksheet":
    for name in wb.sheetnames:
        upper = name.upper()
        if any(hint in upper for hint in hints):
            return wb[name]
    raise BRDFormatError(
        f"Could not find a sheet matching {hints} among {wb.sheetnames}"
    )


def _col_a_values(ws):
    """Yield (row_idx, value_of_col_A) for every non-empty cell in column A."""
    for row in range(1, ws.max_row + 1):
        val = ws.cell(row=row, column=1).value
        if val is not None:
            yield row, str(val).strip()


def _find_label_row(ws, label: str, start: int = 1) -> int:
    """Row number whose column-A text matches `label` (case-insensitive)."""
    for row, val in _col_a_values(ws):
        if row < start:
            continue
        if val.strip().upper() == label.upper():
            return row
    raise BRDFormatError(f"Label '{label}' not found in sheet '{ws.title}'")


def _value_for_label(ws, label: str, start: int = 1) -> str:
    row = _find_label_row(ws, label, start=start)
    val = ws.cell(row=row, column=2).value
    return "" if val is None else str(val).strip()


# --------------------------------------------------------------------------- #
# TECH DOC parsing: header fields + COLUMN NAME / DATA TYPE table
# --------------------------------------------------------------------------- #

def _parse_tech_doc(ws) -> tuple[dict, list[tuple[str, str]]]:
    header = {
        "source_path": _value_for_label(ws, "SOURCE LOCATION (STAGING)"),
        "target_location": _value_for_label(ws, "TARGET LOCATION (RAW_TRUSTED)"),
        "region": _value_for_label(ws, "REGION"),
        "market": _value_for_label(ws, "MARKET"),
        "domain": _value_for_label(ws, "DOMAIN"),
        "source_name": _value_for_label(ws, "SOURCE"),
        "table_name": _value_for_label(ws, "TABLE NAME"),
        "cadence": _value_for_label(ws, "CADENCE"),
        "load_type": _value_for_label(ws, "LOAD TYPE"),
        "primary_key_raw": _value_for_label(ws, "PRIMARY KEY"),
    }

    col_name_row = _find_label_row(ws, "COLUMN NAME")
    columns: list[tuple[str, str]] = []
    row = col_name_row + 1
    while row <= ws.max_row:
        name = ws.cell(row=row, column=1).value
        dtype = ws.cell(row=row, column=2).value
        if name is None or str(name).strip() == "":
            break
        columns.append((str(name).strip(), "" if dtype is None else str(dtype).strip()))
        row += 1

    if not columns:
        raise BRDFormatError("No rows found under COLUMN NAME / DATA TYPE in TECH DOC sheet")

    return header, columns


# --------------------------------------------------------------------------- #
# STTM parsing: SOURCE COLUMN NAME -> TARGET COLUMN NAME (STANDARDIZED)
# --------------------------------------------------------------------------- #

def _parse_sttm(ws) -> dict[str, str]:
    header_row = _find_label_row(ws, "SOURCE COLUMN NAME")
    mapping: dict[str, str] = {}
    row = header_row + 1
    while row <= ws.max_row:
        src = ws.cell(row=row, column=1).value
        tgt = ws.cell(row=row, column=2).value
        if src is None or str(src).strip() == "":
            break
        mapping[str(src).strip()] = "" if tgt is None else str(tgt).strip()
        row += 1
    return mapping


# --------------------------------------------------------------------------- #
# Type / naming normalization
# --------------------------------------------------------------------------- #

def _normalize_sql_type(raw_type: str, warnings: list) -> str:
    base = re.sub(r"\(.*\)", "", raw_type).strip().lower()
    if base in _STRING_TYPES:
        return "string"
    if base in _DATE_TYPES:
        return "date"
    if base in _TIMESTAMP_TYPES:
        return "timestamp"
    if base in _DOUBLE_TYPES:
        return "double"
    if base == "boolean" or base == "bool":
        warnings.append(
            f"Column data type '{raw_type}' (BOOLEAN) has no MDIF equivalent in "
            "utils/schema_config.py — defaulted to 'string', please review."
        )
        return "string"
    warnings.append(
        f"Unrecognized SQL type '{raw_type}' — defaulted to 'string', please review."
    )
    return "string"


def _slug(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")


def _region_code(region: str) -> str:
    key = region.strip().upper()
    if key in _REGION_CODES:
        return _REGION_CODES[key]
    # fallback: acronym from initials, e.g. "SOUTH EAST ASIA" -> "sea"
    return "".join(w[0] for w in key.split() if w).lower() or _slug(region)


def _resolve_primary_key(primary_key_raw: str, src_to_tgt: dict[str, str], warnings: list) -> list:
    if not primary_key_raw or primary_key_raw.strip().upper() in {
        "NO PRIMARY KEY", "NONE", "N/A", "NA", "-",
    }:
        return []
    parts = re.split(r"[,/]", primary_key_raw)
    resolved = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if part in src_to_tgt:
            resolved.append(src_to_tgt[part])
        else:
            # Try case-insensitive match against source names, else fall back
            # to slugifying the raw value so the config isn't silently dropped.
            match = next(
                (v for k, v in src_to_tgt.items() if k.lower() == part.lower()), None
            )
            if match:
                resolved.append(match)
            else:
                warnings.append(
                    f"Primary key column '{part}' not found in STTM mapping — "
                    f"used best-effort slug '{_slug(part)}', please verify."
                )
                resolved.append(_slug(part))
    return resolved


def _detect_pii(column_names: list) -> list:
    hits = []
    for col in column_names:
        low = col.lower()
        if any(hint in low for hint in _PII_HINTS):
            hits.append(col)
    return hits


# --------------------------------------------------------------------------- #
# Main per-file generation
# --------------------------------------------------------------------------- #

def generate_one(
    brd_path: Path,
    repo_root: Path,
    notify_emails: list[str],
    default_file_format: str,
) -> BRDResult:
    warnings: list = []

    wb = openpyxl.load_workbook(brd_path, data_only=True)
    tech_ws = _find_sheet(wb, _TECH_DOC_SHEET_HINTS)
    sttm_ws = _find_sheet(wb, _STTM_SHEET_HINTS)

    header, tech_columns = _parse_tech_doc(tech_ws)
    src_to_tgt = _parse_sttm(sttm_ws)

    if not header["region"] or not header["market"] or not header["source_name"] or not header["table_name"]:
        raise BRDFormatError(
            f"{brd_path.name}: missing REGION/MARKET/SOURCE/TABLE NAME in TECH DOC sheet"
        )

    # Build schema: standardized (STTM) column name + normalized data type, in
    # the same order the columns appear in the TECH DOC sheet.
    schema = []
    target_names = []
    for src_name, raw_type in tech_columns:
        target = src_to_tgt.get(src_name)
        if not target:
            # case-insensitive fallback, else slugify the source name
            target = next(
                (v for k, v in src_to_tgt.items() if k.lower() == src_name.lower()), None
            )
        if not target:
            warnings.append(
                f"Source column '{src_name}' has no STTM mapping — used a "
                f"slugified fallback name, please verify against the STTM tab."
            )
            target = _slug(src_name)
        data_type = _normalize_sql_type(raw_type, warnings)
        schema.append({"column_name": target.lower(), "data_type": data_type})
        target_names.append(target)

    if len(src_to_tgt) != len(tech_columns):
        warnings.append(
            f"TECH DOC lists {len(tech_columns)} column(s) but STTM lists "
            f"{len(src_to_tgt)} mapping(s) — double check both tabs are in sync."
        )

    region = header["region"]
    market = header["market"]
    domain = header["domain"]
    source_name = header["source_name"]
    table_name = _slug(header["table_name"])
    cadence = header["cadence"]
    load_type = _slug(header["load_type"]) or "append"

    job_name = f"{_slug(source_name)}_{_region_code(region)}_{_slug(market)}_{_slug(cadence)}"

    primary_key = _resolve_primary_key(header["primary_key_raw"], src_to_tgt, warnings)
    pii_candidates = _detect_pii(target_names)

    usecols = f"A:{get_column_letter(len(schema))}"

    landing: dict = {
        "source_path": header["source_path"],
        "file_format": default_file_format,
        "pii_columns": pii_candidates,
    }
    if default_file_format in ("excel", "csv", "json"):
        opts_key = f"{default_file_format}_options"
        if default_file_format == "excel":
            landing[opts_key] = {
                "sheet_index": 0,
                "header_row": 1,
                "data_start_row": 2,
                "data_end_row": 1048576,
                "usecols": usecols,
            }
        elif default_file_format == "csv":
            landing[opts_key] = {"delimiter": ","}
        else:  # json
            landing[opts_key] = {}
        warnings.append(
            f"file_format defaulted to '{default_file_format}' — the BRD doesn't "
            "state a landing format explicitly, please confirm against the "
            "actual source file before merging."
        )

    if not header["source_path"]:
        warnings.append("SOURCE LOCATION (STAGING) was blank in the TECH DOC sheet.")

    if not notify_emails:
        notify_emails = ["REPLACE_ME@company.com"]
        warnings.append(
            "No notify emails were provided to the workflow — placeholder "
            "'REPLACE_ME@company.com' written, please replace before merging."
        )

    config = {
        "region": region,
        "market": market,
        "domain": domain,
        "source_name": source_name,
        "notify": {"emails": notify_emails},
        "trigger": {"type": ""},
        "ingestion_steps": {
            "landing": landing,
            "raw": {
                "schema_config": "enforce",
                "deduplication": True,
            },
            "raw_trusted": {
                "load_type": load_type,
                "primary_key": primary_key,
                "enable_expiration": False,
                "validations": ["row_count_validation", "schema_validation"],
                "schema_evolution": True,
            },
        },
        "schema": schema,
    }

    out_dir = repo_root / "metadata" / "configs" / job_name
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{table_name}.json"
    out_path.write_text(json.dumps(config, indent=2) + "\n")

    return BRDResult(
        brd_file=str(brd_path),
        job_name=job_name,
        table_name=table_name,
        output_path=str(out_path.relative_to(repo_root)),
        region=region,
        market=market,
        domain=domain,
        source_name=source_name,
        cadence=cadence,
        load_type=load_type,
        primary_key=primary_key,
        column_count=len(schema),
        file_format=default_file_format,
        pii_candidates=pii_candidates,
        warnings=warnings,
    )


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("brd_files", nargs="+", help="Path(s) to BRD .xlsx workbook(s)")
    parser.add_argument("--repo-root", default=".", help="Repository root (default: cwd)")
    parser.add_argument(
        "--notify-emails", default="",
        help="Comma-separated list of emails for notify.emails",
    )
    parser.add_argument(
        "--file-format", default="excel", choices=["excel", "csv", "parquet", "json", "delta"],
        help="Landing file_format to assume (BRDs don't state this explicitly)",
    )
    parser.add_argument(
        "--summary-out", default="",
        help="If set, write a JSON summary of all generated configs here",
    )
    args = parser.parse_args(argv)

    repo_root = Path(args.repo_root).resolve()
    notify_emails = [e.strip() for e in args.notify_emails.split(",") if e.strip()]

    results = []
    had_error = False
    for f in args.brd_files:
        brd_path = Path(f).resolve()
        try:
            result = generate_one(brd_path, repo_root, notify_emails, args.file_format)
            results.append(result)
            print(f"OK   {brd_path.name} -> {result.output_path}")
            for w in result.warnings:
                print(f"     WARNING: {w}")
        except BRDFormatError as e:
            had_error = True
            print(f"FAIL {brd_path.name}: {e}", file=sys.stderr)

    if args.summary_out:
        Path(args.summary_out).write_text(
            json.dumps([r.__dict__ for r in results], indent=2)
        )

    if had_error:
        return 1
    if not results:
        print("No BRD files were processed.", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
